# -*- coding: utf-8 -*-
import statistics
from collections import Counter, OrderedDict
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analysis import (P, N, count, crosstab, FAMILY_ORDER, INPUT_ORDER, OUTPUT_ORDER, VAL_ORDER, BINS)
from content import MODEL_TEXT, GAPS, BENCHMARKS, RECOMMENDATION
from removed_data import REMOVED, EXTRA_CANDIDATES
import os

OUT = os.path.join(os.path.dirname(__file__), "..", "DPHE_ML_Literature_Log_v2_Verified.xlsx")

NAVY = "1F3864"
FILL_T1 = PatternFill("solid", fgColor="C6EFCE")
FILL_T2 = PatternFill("solid", fgColor="DDEBF7")
FILL_T3 = PatternFill("solid", fgColor="E7E6E6")
FILL_BM = PatternFill("solid", fgColor="FFFF00")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_KPI = PatternFill("solid", fgColor="F2F2F2")
FONT_HDR = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TIERFILL = {"Tier 1": FILL_T1, "Tier 2": FILL_T2, "Tier 3": FILL_T3}


def header(ws, row, cols, widths=None, height=45):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.fill, cell.font, cell.border = FILL_HDR, FONT_HDR, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def write_rows(ws, start, rows, fills=None, height=None):
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            if isinstance(v, (list, tuple)):
                v = "; ".join(v)
            c = ws.cell(row=start + i, column=j, value=v)
            c.alignment, c.border = WRAP, BORDER
            c.font = Font(size=9)
            if fills and fills[i] is not None:
                c.fill = fills[i]
        if height:
            ws.row_dimensions[start + i].height = height


def yn(b):
    return "Yes" if b else "No"


wb = Workbook()

# ------------------------------------------------------------------ 1. Literature Log
ws = wb.active
ws.title = "Literature Log"
LOG_COLS = [
    ("Paper ID", "id", 8), ("Tier", "tier", 8), ("Previous tier (first pass)", "old_tier", 12),
    ("Benchmark Status", "benchmark_status", 10), ("Benchmark Role", "benchmark_role", 14),
    ("PDF Verification Status", "pdf_status", 14), ("Evidence Level", "evidence", 8),
    ("Relevance Score (1-5)", "relevance_score", 9), ("Methodological Quality Score (1-5)", "quality_score", 11),
    ("Removal Reason", "removal_reason", 10),
    ("Full title", "title", 40), ("Authors", "authors", 30), ("Year", "year", 6), ("Journal", "journal", 26),
    ("DOI", "doi", 24), ("Journal quartile", "quartile", 10), ("Quartile Source", "q_source", 24), ("Quartile Year", "q_year", 10),
    ("Original PDF filename", "pdf_original", 26), ("Renamed PDF filename", "pdf_new", 34),
    ("Heat exchanger type", "hx_type", 34), ("HX class", "hx_class", 14), ("Data source", "data_source", 18),
    ("Working fluid", "working_fluid", 26), ("Fluid class", "fluid_class", 12), ("Nanofluid type", "nf_type", 14),
    ("Nanoparticle material", "np_material", 16), ("Base fluid", "base_fluid", 12), ("Concentration range", "conc_range", 16),
    ("Hot-side inlet temperature", "hot_T", 14), ("Cold-side inlet temperature", "cold_T", 14),
    ("Hot-side flow", "hot_flow", 14), ("Cold-side flow", "cold_flow", 14), ("Reynolds-number range", "re_range", 18),
    ("Number of experimental operating conditions", "n_conditions", 22), ("Number of ML samples", "n_samples", 20),
    ("Repeated readings treated as separate samples?", "repeats_as_samples", 16),
    ("ML model(s) tested", "models", 34), ("Model families (standardised)", "families", 26), ("Best-performing model", "best_model", 18),
    ("Input variables", "inputs", 30), ("Input categories", "input_cats", 22), ("Output / target variables", "outputs", 26),
    ("Thermal output categories", "output_cats", 20), ("Data preprocessing", "preprocessing", 18), ("Scaling / normalisation", "scaling", 18),
    ("Feature engineering", "feature_eng", 18), ("Dimensionless inputs?", "dimless", 10), ("Physics-based features", "physics_features", 16),
    ("Correlation-anchored?", "corr_anchored", 16), ("Train/validation/test strategy", "split", 30), ("Random split?", "random_split", 10),
    ("Grouped split?", "grouped", 10), ("k-fold?", "kfold", 12), ("Nested CV?", "nested", 8), ("Leave-one-condition-out?", "loco", 10),
    ("Leave-one-concentration-out?", "loco_conc", 10), ("Leave-one-Re-level-out?", "lo_re", 10), ("External validation?", "external", 12),
    ("Independent confirmation experiments?", "confirm", 12), ("Hyperparameter optimisation", "hpo", 26), ("Model architecture", "architecture", 24),
    ("ANN layer structure", "ann_layers", 14), ("Activation functions", "activation", 12), ("Optimizer", "optimizer", 14), ("Epochs", "epochs", 12),
    ("Regularisation", "regularization", 14), ("Feature importance", "feat_importance", 16), ("SHAP", "shap", 8), ("Sensitivity analysis", "sensitivity", 10),
    ("Applicability-domain checking", "ad_check", 14), ("Uncertainty analysis (predictive)", "uq", 16), ("Prediction intervals", "pi", 14),
    ("Extrapolation testing", "extrap_test", 14), ("Interpolation testing", "interp_test", 16), ("Software / platform", "software", 18),
    ("R2", "r2", 26), ("R", "r", 20), ("RMSE", "rmse", 22), ("MAE", "mae", 18), ("MAPE", "mape", 18), ("MSE", "mse", 18),
    ("Maximum error", "maxerr", 14), ("% within +-5%", "within5", 8), ("% within +-10%", "within10", 8), ("Other reported metrics", "other_metrics", 26),
    ("Metric units (scaled vs physical)", "metric_units", 14), ("Are test metrics truly independent?", "test_independent", 20),
    ("Experimental measurement uncertainty", "exp_uncertainty", 16),
    ("Main methodological strength", "strength", 34), ("Main methodological weakness", "weakness", 34), ("Data leakage risk", "leakage_risk", 18),
    ("Overfitting risk", "overfit_risk", 14), ("Dataset-size concern", "dataset_concern", 16), ("Model capacity excessive for sample size?", "capacity_excess", 14),
    ("Is reported accuracy credible?", "credible", 16), ("Are extrapolation claims justified?", "extrap_justified", 14), ("Is uncertainty addressed?", "uq_addressed", 14),
    ("Would the model realistically reduce experiments?", "reduce_experiments", 18), ("Direct relevance to our project", "relevance", 24),
    ("Methodological element worth adopting", "adopt", 30), ("Methodological element to avoid", "avoid", 30),
    ("Validation category (analysis)", "validation_cat", 16), ("Dataset-size bin (analysis)", "size_bin", 10),
]
header(ws, 1, [c[0] for c in LOG_COLS], [c[2] for c in LOG_COLS], height=60)
rows, fills = [], []
for d in P:
    r = []
    for _, k, _ in LOG_COLS:
        v = d.get(k, "")
        if v is None:
            v = "n/a (abstract only)" if k == "quality_score" else ""
        r.append(v)
    rows.append(r)
    fills.append(TIERFILL[d["tier"]])
write_rows(ws, 2, rows, fills, height=120)
for i, d in enumerate(P):
    if d["benchmark_status"] == "Yes":
        for j in range(1, len(LOG_COLS) + 1):
            ws.cell(row=2 + i, column=j).fill = FILL_BM
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(LOG_COLS))}{N + 1}"

# ------------------------------------------------------------------ 2. Core ML Evidence
ws = wb.create_sheet("Core ML Evidence")
CORE = [("Paper ID", "id", 8), ("Tier", "tier", 8), ("Evidence Level", "evidence", 8), ("Benchmark?", "benchmark_role", 14),
        ("HX type", "hx_class", 16), ("Fluid type", "fluid_class", 13), ("Nanofluid type / material", None, 18),
        ("Dataset size (ML samples)", "n_samples", 18), ("Experimental conditions", "n_conditions", 20),
        ("Repeats treated as samples?", "repeats_as_samples", 14), ("Inputs", "inputs", 28), ("Thermal outputs", "output_cats", 20),
        ("ML algorithms", "families", 24), ("Best model", "best_model", 16), ("Validation method", "split", 28),
        ("Main metrics", None, 30), ("Uncertainty reported?", "uq_flag", 10), ("Physics-informed / anchored?", "phys_flag", 10),
        ("Main weakness", "weakness", 34), ("What we should adopt", "adopt", 32), ("What we should avoid", "avoid", 32),
        ("PDF status", "pdf_status", 14)]
header(ws, 1, [c[0] for c in CORE], [c[2] for c in CORE], height=45)
rows, fills = [], []
for d in P:
    r = []
    for name, k, _ in CORE:
        if name.startswith("Nanofluid type"):
            r.append(f"{d['nf_type']} - {d['np_material']}")
        elif name == "Main metrics":
            parts = [x for x in (d["r2"], d["r"], d["rmse"], d["mape"], d["mae"]) if x and x != "Not reported"]
            r.append(" | ".join(parts[:3]) if parts else "Not reported")
        elif name == "Benchmark?":
            r.append(d["benchmark_role"] or "No")
        else:
            r.append(d.get(k, ""))
    rows.append(r)
    fills.append(FILL_BM if d["benchmark_status"] == "Yes" else TIERFILL[d["tier"]])
write_rows(ws, 2, rows, fills, height=110)
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(CORE))}{N + 1}"

# ------------------------------------------------------------------ 3. Model Comparison
ws = wb.create_sheet("Model Comparison")
FAMMAP = OrderedDict([
    ("Correlation / power-law baseline", None),
    ("Linear / polynomial regression", ["Linear regression", "Polynomial regression"]),
    ("Ridge / Lasso / Elastic Net", ["Regularized linear (Ridge/Lasso/ElasticNet)"]),
    ("ANN / MLP", ["ANN/MLP"]), ("RBF network / GRNN", ["RBF network"]), ("ANFIS", ["ANFIS"]), ("SVR", ["SVR"]), ("GPR", ["GPR"]),
    ("Random Forest", ["Random Forest", "Decision tree"]),
    ("Gradient Boosting / XGBoost / CatBoost", ["Gradient Boosting", "XGBoost", "CatBoost", "AdaBoost"]),
    ("Deep learning", ["Deep learning"]), ("Physics-informed / anchored hybrids", ["Physics-informed / hybrid"]),
])
CORR_PAPERS = {"P13", "P14", "P15", "P27", "P29"}
cols = ["Algorithm / family", "Number of papers using it", "Paper IDs", "Tier 1 papers using it", "Experimental-data papers using it",
        "Nanofluid / hybrid papers using it", "Typical dataset size (median of reported n)", "Typical inputs", "Typical thermal outputs",
        "Typical reported performance", "Strengths", "Weaknesses", "Small-data suitability", "Extrapolation ability",
        "Interpretability", "Uncertainty capability", "Computational complexity", "Relevance to our project"]
header(ws, 1, cols, [26, 10, 22, 10, 12, 12, 12, 24, 24, 30, 30, 30, 14, 16, 12, 16, 12, 28], height=60)
rows = []
for name, fams in FAMMAP.items():
    if fams is None:
        sub = [d for d in P if d["id"] in CORR_PAPERS]
    else:
        sub = [d for d in P if any(f in d["families"] for f in fams)]
    ns = [d["n_numeric"] for d in sub if d["n_numeric"]]
    ic = Counter(x for d in sub for x in d["input_cats"]).most_common(3)
    oc = Counter(x for d in sub for x in d["output_cats"]).most_common(3)
    t = MODEL_TEXT[name]
    rows.append([name, len(sub), ", ".join(d["id"] for d in sub), sum(d["tier"] == "Tier 1" for d in sub),
                 sum(d["src_class"] in ("Experimental", "Mixed") for d in sub),
                 sum(d["fluid_class"] != "Conventional fluid" for d in sub),
                 (f"{int(statistics.median(ns))} (n={len(ns)} papers)" if ns else "Not reported"),
                 ", ".join(f"{k} ({v})" for k, v in ic) or "-", ", ".join(f"{k} ({v})" for k, v in oc) or "-",
                 t["typical_perf"], t["strengths"], t["weaknesses"], t["small"], t["extrap"], t["interp"], t["uq"], t["cost"], t["rel"]])
write_rows(ws, 2, rows, height=95)
r0 = len(rows) + 3
ws.cell(row=r0, column=1, value="Notes").font = Font(bold=True)
notes = ["Counts are computed from the 24 retained papers in the Literature Log (multi-label: a paper is counted once per family it tests).",
         "'Correlation / power-law baseline' counts papers that compare ML against an empirical correlation (P13, P14, P15, P27, P29); P27 additionally anchors the learner to the correlation.",
         "'Random Forest' includes the single decision-tree model (P16); 'Gradient Boosting' includes BRT, HistGBR, XGBoost, CatBoost and AdaBoost.",
         "Performance figures are as reported by the authors and are mostly from random splits; they are not comparable across papers and overstate generalisation (see Validation critique).",
         "No retained paper used CatBoost except P04; no retained paper used Bayesian neural networks or conformal prediction."]
for i, n in enumerate(notes):
    c = ws.cell(row=r0 + 1 + i, column=1, value=n)
    ws.merge_cells(start_row=r0 + 1 + i, start_column=1, end_row=r0 + 1 + i, end_column=12)
    c.alignment = WRAP
ws.freeze_panes = "B2"

# ------------------------------------------------------------------ 4. Research Gaps
ws = wb.create_sheet("Research Gaps")
header(ws, 1, ["#", "Gap", "Supporting papers", "Evidence Level supporting the claim", "Why it matters", "Opportunity for our study", "Priority", "Strength of evidence"],
       [4, 44, 40, 18, 36, 36, 10, 22], height=40)
write_rows(ws, 2, [[i + 1, g["gap"], g["papers"], g["evidence"], g["why"], g["opp"], g["priority"], g["strength"]] for i, g in enumerate(GAPS)], height=95)
r0 = len(GAPS) + 3
ws.cell(row=r0, column=2, value="Wording rule: every gap is stated relative to the reviewed (retained) literature. Absence in this set is not proof that no such study exists anywhere.").alignment = WRAP
ws.cell(row=r0 + 1, column=2, value="Previous-review gap removed after full-text reading: 'Hydraulic penalty is left out of the model' - out of the thermal-only scope of the project and therefore not a target gap.").alignment = WRAP
ws.cell(row=r0 + 2, column=2, value="Previous-review gap re-worded: 'No ML study exists for MgO-TiO2 hybrid nanofluid in a DPHE' - cannot be verified from this database; replaced by Gap 1 (hybrid nanofluid DPHE ML with verified full text).").alignment = WRAP
ws.freeze_panes = "C2"

# ------------------------------------------------------------------ 5. Statistical Analysis
ws = wb.create_sheet("Statistical Analysis")
for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", [34, 9, 9, 3, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12]):
    ws.column_dimensions[col].width = w
t = ws.cell(row=1, column=1, value="Statistical Literature Dashboard - verified retained database (n = %d papers)" % N)
t.font = Font(bold=True, size=16, color=NAVY)
ws.cell(row=2, column=1, value="Generated from the Literature Log (second-stage full-text verification, 5 Oct 2026). Multi-label analyses count a paper once per category; percentages use n = %d unless stated." % N).font = Font(italic=True, size=9)

def k(cond):
    return sum(1 for d in P if cond(d))

KPIS = [
    ("Total retained papers", N), ("Tier 1", k(lambda d: d["tier"] == "Tier 1")), ("Tier 2", k(lambda d: d["tier"] == "Tier 2")),
    ("Tier 3", k(lambda d: d["tier"] == "Tier 3")), ("Evidence A", k(lambda d: d["evidence"] == "A")), ("Evidence B", k(lambda d: d["evidence"] == "B")),
    ("Evidence C", k(lambda d: d["evidence"] == "C")), ("Experimental", k(lambda d: d["src_class"] == "Experimental")),
    ("CFD", k(lambda d: d["src_class"] == "CFD")), ("Analytical", k(lambda d: d["src_class"] == "Analytical")), ("Mixed", k(lambda d: d["src_class"] == "Mixed")),
    ("Nanofluid (mono)", k(lambda d: d["fluid_class"] == "Nanofluid")), ("Hybrid nanofluid", k(lambda d: d["fluid_class"] == "Hybrid nanofluid")),
    ("Conventional fluid", k(lambda d: d["fluid_class"] == "Conventional fluid")), ("Benchmark papers", k(lambda d: d["benchmark_status"] == "Yes")),
    ("With predictive UQ", k(lambda d: d["uq_flag"] == "Yes")), ("With independent validation", k(lambda d: d["indep_flag"] == "Yes")),
    ("Physics-informed / anchored", k(lambda d: d["phys_flag"] == "Yes")),
]
r = 4
for i, (lab, val) in enumerate(KPIS):
    rr, cc = r + (i // 6) * 3, 1 + (i % 6) * 2 if i % 6 else 1
    cc = [1, 2, 5, 7, 9, 11][i % 6]
    a = ws.cell(row=rr, column=cc, value=lab)
    b = ws.cell(row=rr + 1, column=cc, value=val)
    for c in (a, b):
        c.fill, c.border = FILL_KPI, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    a.font = Font(size=9, bold=True, color="595959")
    b.font = Font(size=20, bold=True, color=NAVY)
    ws.row_dimensions[rr + 1].height = 30
row = r + 10

charts_meta = []


def block(title, od, note=None, pct=True, horizontal=True, sort=True, denom=N):
    global row
    items = list(od.items())
    if sort:
        items = sorted(items, key=lambda kv: -kv[1])
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color=NAVY)
    if note:
        ws.cell(row=row + 1, column=1, value=note).font = Font(italic=True, size=8)
    h = row + 2
    for j, c in enumerate(["Category", "Count", "%"]):
        cell = ws.cell(row=h, column=1 + j, value=c)
        cell.fill, cell.font, cell.border = FILL_SUB, Font(bold=True, size=9), BORDER
    for i, (kk, v) in enumerate(items):
        ws.cell(row=h + 1 + i, column=1, value=str(kk)).border = BORDER
        ws.cell(row=h + 1 + i, column=2, value=v).border = BORDER
        pc = ws.cell(row=h + 1 + i, column=3, value=(v / denom) if pct else None)
        pc.number_format = "0%"
        pc.border = BORDER
        for j in range(1, 4):
            ws.cell(row=h + 1 + i, column=j).font = Font(size=9)
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.style = 10
    ch.title = title
    ch.y_axis.title = "Number of papers"
    ch.legend = None
    data = Reference(ws, min_col=2, min_row=h, max_row=h + len(items))
    cats = Reference(ws, min_col=1, min_row=h + 1, max_row=h + len(items))
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.height = max(6.5, 0.75 * len(items) + 3)
    ch.width = 17
    ch.y_axis.majorGridlines = None
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
    ch.series[0].graphicalProperties.solidFill = "2F5597"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, f"E{row}")
    span = max(len(items) + 4, int(ch.height / 0.5) + 2)
    charts_meta.append((title, items))
    row += span + 1


def xtab(title, t, note=None):
    global row
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color=NAVY)
    if note:
        ws.cell(row=row + 1, column=1, value=note).font = Font(italic=True, size=8)
    h = row + 2
    corder = list(next(iter(t.values())).keys())
    ws.cell(row=h, column=1, value="").fill = FILL_SUB
    for j, c in enumerate(corder):
        cell = ws.cell(row=h, column=2 + j if j == 0 else 4 + j, value=c)
    # place columns: B then E.. to keep column C free for readability
    cols_idx = [2] + [5 + j for j in range(len(corder) - 1)]
    for j, c in enumerate(corder):
        cell = ws.cell(row=h, column=cols_idx[j], value=c)
        cell.fill, cell.font, cell.border = FILL_SUB, Font(bold=True, size=8), BORDER
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    ws.row_dimensions[h].height = 42
    for i, (rk, rv) in enumerate(t.items()):
        a = ws.cell(row=h + 1 + i, column=1, value=rk)
        a.font, a.border = Font(size=9, bold=True), BORDER
        for j, c in enumerate(corder):
            cell = ws.cell(row=h + 1 + i, column=cols_idx[j], value=rv[c])
            cell.border, cell.font = BORDER, Font(size=9)
            cell.alignment = Alignment(horizontal="center")
    rng = f"{get_column_letter(cols_idx[0])}{h + 1}:{get_column_letter(cols_idx[-1])}{h + len(t)}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="5B9BD5"))
    row = h + len(t) + 3


block("A. Papers by publication year", count("year", [2022, 2023, 2024, 2025, 2026]), sort=False, horizontal=False)
block("B. Papers by Tier", count("tier", ["Tier 1", "Tier 2", "Tier 3"]), sort=False, horizontal=False)
block("C. Papers by Evidence Level", count("evidence", ["A", "B", "C"]), "A = full paper read; B = full PDF, key sections verified by targeted extraction; C = abstract only", sort=False, horizontal=False)
block("D. Papers by journal quartile", count("quartile"), "SCImago could not be reached from this session (network policy); only P01/P09 carry a quartile verified in the first-pass session.")
block("E. Papers by heat-exchanger type", count("hx_class"))
block("F. Papers by data source", count("src_class", ["Experimental", "CFD", "Analytical", "Mixed"]))
block("G. Papers by fluid type", count("fluid_class", ["Conventional fluid", "Nanofluid", "Hybrid nanofluid"]))
block("H. Papers by nanoparticle material", count("np_group"))
block("I. Frequency of ML model families", count("families", FAMILY_ORDER, multi=True), "Multi-label; a paper is counted once for each family it tests.")
block("J. Single-model vs multi-model comparison studies", OrderedDict([("Multi-model comparison", k(lambda d: d["multi_model"] == "Yes")), ("Single model", k(lambda d: d["multi_model"] == "No"))]), horizontal=False)
block("K. Frequency of input-variable categories", count("input_cats", INPUT_ORDER, multi=True), "Multi-label. Note: no study used cold inlet temperature as an input.")
block("L. Frequency of predicted thermal outputs", count("output_cats", OUTPUT_ORDER, multi=True), "Multi-label; hydraulic outputs deliberately excluded (thermal-only scope). No study predicted outlet temperatures in the retained set.")
block("M. Validation-method frequency", count("val_multi", VAL_ORDER, multi=True), "Multi-label. 'Grouped' and 'leave-one-group-out' occur only outside double-pipe heat exchangers (P27, P29).", sort=False)
block("N. Dataset-size distribution (ML samples)", count("size_bin", BINS), "Bins of the reported number of ML samples. Several small studies do not report n.", sort=False, horizontal=False)
block("O. Predictive uncertainty reported", count("uq_flag", ["Yes", "No"]), sort=False, horizontal=False)
block("P. Physics-informed / correlation-anchored features", count("phys_flag", ["Yes", "No"]), sort=False, horizontal=False)
block("Q. Dimensionless input variables", count("dimless_flag", ["Yes", "Partial", "No / not reported"]), sort=False, horizontal=False)
REP = OrderedDict([("Dataset size", "dataset"), ("Exact train/test split", "split"), ("Hyperparameters", "hyper"), ("Scaling method", "scaling"),
                   ("Exact input features", "inputs"), ("Exact output variables", "outputs"), ("Predictive uncertainty", "uncertainty"),
                   ("Independent validation", "indep"), ("Extrapolation testing", "extrap"), ("Applicability-domain checking", "ad")])
block("R. Reporting-quality analysis (papers reporting each item)", OrderedDict((lab, k(lambda d, kk=kk: d["reporting"][kk])) for lab, kk in REP.items()), sort=False)

ws.cell(row=row, column=1, value="CROSS-TAB ANALYSES").font = Font(bold=True, size=14, color=NAVY)
row += 2
FAM_SHORT = ["ANN/MLP", "ANFIS", "SVR", "GPR", "Random Forest", "Gradient Boosting", "XGBoost", "CatBoost", "Linear regression", "Polynomial regression",
             "Regularized linear (Ridge/Lasso/ElasticNet)", "Deep learning", "Physics-informed / hybrid"]
xtab("X1. ML model family x fluid type", crosstab("families", "fluid_class", FAM_SHORT, ["Conventional fluid", "Nanofluid", "Hybrid nanofluid"], rmulti=True))
xtab("X2. ML model family x heat-exchanger class", crosstab("families", "hx_class", FAM_SHORT, ["Double-pipe / concentric", "Tube / coil test section", "Other"], rmulti=True))
xtab("X3. ML model family x dataset-size category", crosstab("families", "size_bin", FAM_SHORT, BINS, rmulti=True))
xtab("X4. Validation method x dataset size", crosstab("val_multi", "size_bin", VAL_ORDER, BINS, rmulti=True))
xtab("X5. ML model family x validation quality", crosstab("families", "val_quality", FAM_SHORT, ["None / unclear", "Random only", "Grouped / external"], rmulti=True))
xtab("X6. Thermal target x ML model family", crosstab("output_cats", "families", [o for o in OUTPUT_ORDER if o != "Outlet temperatures"], ["ANN/MLP", "ANFIS", "SVR", "GPR", "Random Forest", "XGBoost", "Gradient Boosting", "Linear regression", "Polynomial regression"], rmulti=True, cmulti=True))
xtab("X7. Data source x headline accuracy level (best R or R2 reported)", crosstab("src_class", "acc_level", ["Experimental", "CFD", "Analytical", "Mixed"], [">=0.99", "0.95-0.99", "<0.95", "R/R2 not reported"]))
xtab("X8. Nanofluid type x thermal output", crosstab("np_group", "output_cats", ["Al2O3", "Graphene family (GO/rGO/GNP)", "MWCNT", "Fe3O4", "SiO2", "Hybrid Al2O3-CuO", "Hybrid Al2O3-TiO2", "Several (property-based, analytical)"], ["Nusselt number", "Heat-transfer coefficient", "Overall heat-transfer coefficient", "Heat-transfer rate", "Effectiveness", "NTU", "Entropy/exergy", "Thermal performance factor"], cmulti=True))
xtab("X9. Tier x data source", crosstab("tier", "src_class", ["Tier 1", "Tier 2", "Tier 3"], ["Experimental", "CFD", "Analytical", "Mixed"]))
xtab("X10. Evidence Level x Tier", crosstab("evidence", "tier", ["A", "B", "C"], ["Tier 1", "Tier 2", "Tier 3"]))
xtab("X11. Multi-model comparison x validation quality", crosstab("multi_model", "val_quality", ["Yes", "No"], ["None / unclear", "Random only", "Grouped / external"]))
xtab("X13. Dataset size x headline accuracy level", crosstab("size_bin", "acc_level", BINS, [">=0.99", "0.95-0.99", "<0.95", "R/R2 not reported"]))
xtab("X14. Dataset size x ML complexity", crosstab("size_bin", "complexity", BINS, ["Low", "Medium", "High", "Unknown"]),
     "Complexity: Low = linear/regularised; Medium = kernel/shallow trees/single-hidden-layer ANN/ANFIS; High = >=2 hidden layers, deep trees, metaheuristic-tuned or deep nets.")
xtab("X15. ANN usage x sample size", crosstab("uses_ann", "size_bin", ["Yes", "No"], BINS))
ws.cell(row=row, column=1, value="X12 (journal quartile x methodological quality) is not shown: quartiles could not be verified in this session, so the cross-tab would be uninformative.").font = Font(italic=True, size=9)
row += 3

STAT_ROW_FINDINGS = row
ws.cell(row=row, column=1, value="KEY STATISTICAL FINDINGS").font = Font(bold=True, size=14, color=NAVY)
row += 1
FINDINGS = []
nann = k(lambda d: "ANN/MLP" in d["families"])
FINDINGS.append(f"ANN/MLP is the dominant model family: {nann} of {N} retained papers ({nann / N:.0%}) use it; GPR appears in only 1 paper (P03) despite most experimental datasets being small.")
nrand = k(lambda d: d["val_quality"] == "Random only")
nnone = k(lambda d: d["val_quality"] == "None / unclear")
t1 = [d for d in P if d["tier"] == "Tier 1"]
t1g = sum(d["val_quality"] == "Grouped / external" for d in t1)
FINDINGS.append(f"Validation is weak: {nrand} papers rely only on random splits and {nnone} report no clear validation; none of the {len(t1)} Tier 1 (double-pipe, experimental) papers used grouped, held-out-condition or external validation (count = {t1g}).")
FINDINGS.append(f"Predictive uncertainty is reported by {k(lambda d: d['uq_flag'] == 'Yes')} of {N} papers (P37 - CFD; P27 - partial), and by none of the double-pipe papers.")
FINDINGS.append(f"Nusselt number is the most frequent thermal target ({sum('Nusselt number' in d['output_cats'] for d in P)} papers), followed by heat-transfer coefficient ({sum('Heat-transfer coefficient' in d['output_cats'] for d in P)}); effectiveness and NTU appear in {sum('Effectiveness' in d['output_cats'] for d in P)} papers each, and no paper learns outlet temperatures.")
hi = [d for d in P if d["acc_level"] == ">=0.99"]
FINDINGS.append(f"Headline R/R2 >= 0.99 is reported by {len(hi)} papers; {sum(d['val_quality'] != 'Grouped / external' for d in hi)} of these rest on random splits or no clear validation, so high R2 is not evidence of generalisation.")
nr = k(lambda d: d["size_bin"] == "Not reported")
FINDINGS.append(f"Dataset size is not reported by {nr} of {N} papers; among Tier 1 papers only P02 (84), P04 (540) and P05 (72) state n.")
FINDINGS.append(f"Physics enters the models rarely: {k(lambda d: d['phys_flag'] == 'Yes')} papers use physics-anchored or physics-structured learning (P27, P28 - both outside heat exchangers), while {k(lambda d: d['dimless_flag'] == 'Yes')} use fully dimensionless inputs.")
FINDINGS.append(f"Hybrid nanofluids appear in {k(lambda d: d['fluid_class'] == 'Hybrid nanofluid')} papers, and only one (P01, abstract only) concerns a double-pipe heat exchanger.")
FINDINGS.append("No retained study used cold-side inlet temperature as an input, and no double-pipe study varied both flow rates and both inlet temperatures together with concentration.")
FINDINGS.append(f"Multi-model comparisons are common ({k(lambda d: d['multi_model'] == 'Yes')} papers), but only P27 and P29 compare families under grouped validation; in P29 the correlation beat ML on unseen geometries (MAE 1,111 vs 1,819 W/m2K).")
for i, f in enumerate(FINDINGS):
    c = ws.cell(row=row + i, column=1, value=f"{i + 1}. {f}")
    ws.merge_cells(start_row=row + i, start_column=1, end_row=row + i, end_column=14)
    c.alignment = WRAP
    c.font = Font(size=10)
    ws.row_dimensions[row + i].height = 32

# ------------------------------------------------------------------ 6. Removed Papers
ws = wb.create_sheet("Removed Papers")
header(ws, 1, ["Original Paper ID", "Title", "Authors", "Year", "Journal", "DOI", "Original Tier", "Removal Reason", "PDF filename", "Notes"],
       [9, 44, 30, 6, 26, 26, 10, 22, 40, 50])
write_rows(ws, 2, [[x["id"], x["title"], x["authors"], x["year"], x["journal"], x["doi"], x["old_tier"], x["reason"], x["pdf"], x["notes"]] for x in REMOVED], height=70)
ws.freeze_panes = "B2"

# ------------------------------------------------------------------ 7. Download / PDF Verification Status
ws = wb.create_sheet("Download PDF Verification")
ws.sheet_properties.tabColor = "70AD47"
header(ws, 1, ["Paper ID", "Title", "PDF available?", "PDF filename (original in Papers folder)", "Renamed filename (target)", "Verified match?",
               "PDF Verification Status", "Evidence Level", "Retained?", "Abstract-only?", "Action needed"],
       [9, 44, 9, 34, 40, 10, 16, 8, 9, 9, 40])
rows = []
by_id = {d["id"]: d for d in P}
rem = {x["id"]: x for x in REMOVED}
for i in range(1, 38):
    pid = f"P{i:02d}"
    if pid in by_id:
        d = by_id[pid]
        has = d["evidence"] in ("A", "B")
        rows.append([pid, d["title"], yn(has), d["pdf_original"], d["pdf_new"], "Yes - title/authors/DOI checked in PDF" if has else "n/a",
                     d["pdf_status"], d["evidence"], "Yes", yn(d["evidence"] == "C"),
                     ("Obtain full text (high priority) and re-assess" if d["evidence"] == "C" else ("Rename file (connector was read-only)"))])
    elif pid in rem:
        x = rem[pid]
        has = "(no PDF)" not in x["pdf"]
        status = "Removed as irrelevant" if has else "PDF unavailable"
        rows.append([pid, x["title"], yn(has), x["pdf"].split(" -> ")[0] if has else "-", x["pdf"].split(" -> ")[1] if has else "-",
                     "Yes" if has else "n/a", status, "A/B" if has else "C", "No", "No" if has else "Yes",
                     ("Rename with REMOVED_ prefix" if has else ("Download (high priority)" if "high priority" in x["notes"] or "ACTION: download." in x["notes"] else "Optional download"))])
    else:
        rows.append([pid, "Not used", "", "", "", "", "", "", "", "", ""])
for x in EXTRA_CANDIDATES:
    rows.append([x[0], x[2], "No", "-", "-", "n/a", "Not logged (screening candidate)", "-", "No", "-", x[5]])
rows.append(["(file)", "Claude-Science-0.1.56.dmg (software installer found in Papers folder)", "n/a", "Claude-Science-0.1.56.dmg", "-", "Not a paper", "-", "-", "-", "-", "Move out of the Papers folder"])
write_rows(ws, 2, rows, height=45)
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:K{len(rows) + 1}"

# ------------------------------------------------------------------ Extra: PDF Rename Map
ws = wb.create_sheet("PDF Rename Map")
header(ws, 1, ["Current filename in Papers folder", "New filename", "Paper ID", "Status"], [58, 60, 9, 12])
pairs = [(d["pdf_original"], d["pdf_new"], d["id"], "Retained") for d in P if d["evidence"] in ("A", "B")]
pairs[[p[2] for p in pairs].index("P37")] = ("Energy Science   Engineering - 2026 - Salah - Inter‐Material Transfer Learning for Accelerated Nanofluid Heat Transfer.pdf",
                                            by_id["P37"]["pdf_new"], "P37", "Retained")
for x in REMOVED:
    if "->" in x["pdf"]:
        a, b = [s.strip() for s in x["pdf"].split("->")]
        pairs.append((a, b, x["id"], "Removed"))
pairs.sort(key=lambda p: p[2])
write_rows(ws, 2, [list(p) for p in pairs], height=18)
ws.cell(row=len(pairs) + 3, column=1, value="The Microsoft 365 connector in this session is read-only (Files.ReadWrite.All not granted), so files could not be renamed in OneDrive. Use rename_pdfs.ps1 (Windows) or rename_pdfs.sh (macOS/Linux) in the Papers folder.").alignment = WRAP

# ------------------------------------------------------------------ Extra: Legend & Method
ws = wb.create_sheet("Legend & Method")
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 110
leg = [
    ("Workbook", "DPHE ML literature database - second-stage verified (full-text) version, 5 Oct 2026. Supersedes DPHE_ML_Literature_Log.xlsx (first pass)."),
    ("Scope", "Thermal-performance prediction only. Hydraulic quantities are recorded when a paper reports them but never drive tiers, gaps or recommendations."),
    ("Tier 1 (green)", "Double-pipe/concentric exchanger + experimental data + ML prediction of a thermal quantity."),
    ("Tier 2 (blue)", "Different geometry or data source but highly transferable methodology (validation, small data, physics-informed learning, uncertainty)."),
    ("Tier 3 (grey)", "Supporting relevance (single aspect, CFD design studies, cautionary examples, frameworks)."),
    ("Benchmark (yellow)", "Selected after full-text review: P02 primary, P13 secondary, P37 supporting (all Evidence A)."),
    ("Evidence A", "Full paper available and read directly (methods, results and conclusions read sequentially)."),
    ("Evidence B", "Full PDF available; key sections (data, ML set-up, validation, results tables, conclusions) verified through targeted extraction; not every paragraph read."),
    ("Evidence C", "Abstract / publisher / index information only."),
    ("PDF Verification Status", "Full PDF verified | Abstract only — retained | PDF unavailable | PDF present but mismatched | Removed as irrelevant. No PDF in the folder was mismatched: all 26 PDFs matched an existing Paper ID by title, authors and DOI read from the PDF."),
    ("'Not reported'", "The item is absent from the source. Values marked '(reviewer inference)' are deductions, not author statements."),
    ("Quartiles", "Not verified in this session: scimagojr.com and api.openalex.org were blocked by the environment network policy. P01/P09 (J. Therm. Anal. Calorim.) keep the Q2 that the first-pass session recorded as verified on SCImago. All others: 'Not verified'. The first-pass labels such as 'Q1 (recalled)' were deleted."),
    ("Relevance Score (1-5)", "5 = same exchanger class, nanofluid, operating inputs and thermal outputs as our rig; 1 = different geometry, data source and purpose."),
    ("Quality Score (1-5)", "5 = leakage-aware grouped/external validation, physical-unit errors, transparent reporting; 1 = no meaningful validation or severe over-parameterisation. 'n/a' for abstract-only papers."),
    ("Leakage taxonomy", "Kapoor & Narayanan (2023, P30 - moved to Removed Papers as supporting reference): L1.1 no test set; L1.2 preprocessing on train+test; L1.3 feature selection on train+test; L1.4 duplicates; L2 illegitimate features; L3.2 non-independence between train and test (e.g. repeated readings / neighbouring factorial points)."),
    ("Interpolation vs extrapolation", "Interpolation = prediction between tested points (random splits test this). Extrapolation = outside the range of any input. True held-out-condition prediction = an operating condition (or concentration / flow level) never seen in training, evaluated with all its repeats withheld."),
    ("Changes vs first pass", "Tiers reassessed (P03 1->2; P12, P15, P16, P18, P21, P22, P26 2->3; P27, P28 3->2; P37 pending->2); benchmarks changed (P03, P04 dropped; P13, P37 added); 13 records moved to Removed Papers; P29 DOI and P33/P37 author lists recovered from PDFs; P37 journal issue/pages added."),
]
for i, (a, b) in enumerate(leg):
    ws.cell(row=i + 1, column=1, value=a).font = Font(bold=True)
    c = ws.cell(row=i + 1, column=2, value=b)
    c.alignment = WRAP
    ws.row_dimensions[i + 1].height = 45
ws.cell(row=1, column=1).fill = FILL_T1

# ------------------------------------------------------------------ Extra: Benchmarks & Recommendation
ws = wb.create_sheet("Benchmarks & Plan")
header(ws, 1, ["Paper", "Role", "Why selected", "Adopt", "Do NOT copy", "What we can improve"], [8, 16, 50, 45, 45, 45])
write_rows(ws, 2, [[b["id"], b["role"], b["why"], b["adopt"], b["not_copy"], b["improve"]] for b in BENCHMARKS], [FILL_BM] * 3, height=140)
r0 = 7
header_row = r0
for j, c in enumerate(["Design element", "Recommendation (thermal-only)"], 1):
    cell = ws.cell(row=header_row, column=j, value=c)
    cell.fill, cell.font = FILL_HDR, FONT_HDR
ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row, end_column=6)
for i, (a, b) in enumerate(RECOMMENDATION):
    ws.cell(row=r0 + 1 + i, column=1, value=a).font = Font(bold=True, size=9)
    c = ws.cell(row=r0 + 1 + i, column=2, value=b)
    ws.merge_cells(start_row=r0 + 1 + i, start_column=2, end_row=r0 + 1 + i, end_column=6)
    c.alignment = WRAP
    c.font = Font(size=9)
    ws.row_dimensions[r0 + 1 + i].height = 48

order = ["Literature Log", "Core ML Evidence", "Model Comparison", "Research Gaps", "Statistical Analysis", "Removed Papers",
         "Download PDF Verification", "Benchmarks & Plan", "PDF Rename Map", "Legend & Method"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
import json
json.dump({"findings": FINDINGS, "charts": [(t, [(str(a), b) for a, b in it]) for t, it in charts_meta]},
          open(os.path.join(os.path.dirname(__file__), "stats.json"), "w"), indent=1)
print("saved", OUT)
